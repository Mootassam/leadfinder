"""
Lead Finder â€” local web app. Run `python server.py` and open the printed URL.
"""

from __future__ import annotations

import csv
import io
import json
import os
import socket
import sqlite3
import sys
import threading
import time
import webbrowser

from flask import Flask, Response, jsonify, request, send_from_directory

import app_paths
import categories
import engine
import overture
import registries
import sources
import store

app = Flask(__name__, static_folder=None)
# store.init() runs in __main__ after the single-instance check: a second launch must not touch
# the database (it would mark the first copy's running searches as stopped)

EXPORT_COLS = [("name", "Company"), ("email", "Email"), ("emails", "All emails"), ("phone", "Phone"),
               ("phones", "All phones"), ("website", "Website"), ("category", "Category"), ("niche", "Niche"),
               ("address", "Address"), ("city", "City"), ("postcode", "Postcode"), ("country", "Country"),
               ("facebook", "Facebook"), ("instagram", "Instagram"), ("linkedin", "LinkedIn"),
               ("twitter", "X / Twitter"), ("youtube", "YouTube"), ("tiktok", "TikTok"),
               ("rating", "Rating"), ("reviews", "Reviews"), ("score", "Lead score"), ("mx_ok", "Email domain OK"),
               ("hours", "Opening hours"), ("lat", "Latitude"), ("lon", "Longitude"), ("maps_url", "Map link"),
               ("legal_name", "Legal name"), ("company_id", "Company number"), ("employees", "Employees"),
               ("founded", "Founded"), ("people", "Decision makers"), ("people_emails", "Decision-maker emails"),
               ("sources", "Sources"), ("note", "Note")]


def _filters(src) -> dict:
    keys = ("search", "q", "has_email", "has_phone", "has_website", "has_social", "has_people", "mx_ok", "starred",
            "not_exported", "min_score", "city", "country", "category", "sort", "ids")
    return {k: src.get(k) for k in keys if src.get(k) not in (None, "", "0", "false")}


@app.get("/")
def index():
    return send_from_directory(app_paths.static_dir(), "index.html")


@app.get("/static/<path:p>")
def static_files(p):
    r = send_from_directory(app_paths.static_dir(), p)
    r.headers["Cache-Control"] = "no-cache"
    return r


# ------------------------------------------------------------------ lookups
@app.get("/api/niches")
def niches():
    return jsonify(categories.catalogue())


@app.get("/api/sources-info")
def sources_info():
    return jsonify({"registries": registries.SUPPORTED, "duckdb": overture.available()})


@app.get("/api/niche-match")
def niche_match():
    m = categories.match(request.args.get("q", ""))
    return jsonify({"label": m["label"], "exact": m["exact"]})


@app.get("/api/places")
def places():
    q = request.args.get("q", "").strip()
    if len(q) < 2:
        return jsonify([])
    try:
        return jsonify(sources.suggest(q))
    except Exception as ex:
        return jsonify({"error": str(ex)}), 502


# ------------------------------------------------------------------ searches / jobs
@app.post("/api/search")
def search():
    d = request.get_json(force=True) or {}
    niches = [x for x in (d.get("niches") or []) if x.strip()][:25]
    locs = [x for x in (d.get("locations") or []) if x.strip()][:50]
    if not niches or not locs:
        return jsonify({"error": "Add at least one business type and one location."}), 400
    params = {"sources": d.get("sources") or {"osm": True, "overture": True, "registry": True},
              "enrich": bool(d.get("enrich", True)), "max_leads": int(d.get("max_leads") or 0),
              "find_websites": bool(d.get("find_websites", True)), "people": bool(d.get("people", True)),
              "enrich_missing_only": bool(d.get("enrich_missing_only", False))}
    sid = engine.start_search(niches, locs, params)
    return jsonify({"id": sid})


@app.get("/api/jobs")
def jobs():
    return jsonify([engine.get(i) for i in engine.running_ids()])


@app.get("/api/jobs/<int:sid>")
def job(sid):
    j = engine.get(sid)
    if not j:
        s = store.get_search(sid)
        if not s:
            return jsonify({"error": "not found"}), 404
        j = {"id": sid, "status": s["status"], "phase": "Finished", "log": s["log"][-80:], "counts": {},
             "recent": [], "params": s["params"], "elapsed": round((s["finished_at"] or s["created_at"]) -
                                                                    s["created_at"], 1)}
    j["stats"] = (store.get_search(sid) or {}).get("stats", {})
    return jsonify(j)


@app.post("/api/jobs/<int:sid>/cancel")
def job_cancel(sid):
    return jsonify({"ok": engine.cancel(sid)})


@app.post("/api/searches/<int:sid>/resume")
def search_resume(sid):
    return jsonify({"ok": engine.resume(sid)})


@app.get("/api/searches")
def searches():
    running = set(engine.running_ids())
    out = store.list_searches()
    for s in out:
        if s["id"] in running:
            s["status"] = "running"
    return jsonify(out)


@app.delete("/api/searches/<int:sid>")
def search_delete(sid):
    engine.cancel(sid)
    store.delete_search(sid, request.args.get("leads") == "1")
    return jsonify({"ok": True})


# ------------------------------------------------------------------ leads
@app.get("/api/leads")
def leads():
    page = max(1, int(request.args.get("page", 1)))
    size = max(10, min(500, int(request.args.get("size", 50))))
    return jsonify(store.query_leads(_filters(request.args), page, size))


@app.get("/api/leads/geo")
def leads_geo():
    return jsonify(store.geo_points(_filters(request.args)))


@app.get("/api/leads/<int:lid>")
def lead(lid):
    d = store.get_lead(lid)
    return (jsonify(d), 200) if d else (jsonify({"error": "not found"}), 404)


@app.patch("/api/leads/<int:lid>")
def lead_update(lid):
    store.update_lead(lid, request.get_json(force=True) or {})
    return jsonify(store.get_lead(lid))


@app.post("/api/leads/delete")
def leads_delete():
    d = request.get_json(force=True) or {}
    ids = _resolve_ids(d)
    store.delete_leads(ids)
    return jsonify({"deleted": len(ids)})


@app.post("/api/leads/recheck")
def leads_recheck():
    d = request.get_json(force=True) or {}
    ids = [r["id"] for r in store.iter_leads(dict(_filters(d.get("filters") or {}), has_website="1"))] \
        if d.get("all") else [int(i) for i in d.get("ids") or []]
    if not ids:
        return jsonify({"error": "No leads with a website selected."}), 400
    return jsonify({"id": engine.enrich_leads(ids[:20000])})


def _resolve_ids(d):
    if d.get("all"):
        return [r["id"] for r in store.iter_leads(_filters(d.get("filters") or {}))]
    return [int(i) for i in d.get("ids") or []]


@app.get("/api/stats")
def stats():
    s = store.stats()
    s["running"] = engine.running_ids()
    return jsonify(s)


# ------------------------------------------------------------------ export
def _flat(l: dict) -> dict:
    out = {}
    for k, _ in EXPORT_COLS:
        if k in ("emails", "phones", "sources"):
            out[k] = "; ".join(l.get(k) or [])
        elif k in ("facebook", "instagram", "linkedin", "twitter", "youtube", "tiktok"):
            out[k] = (l.get("socials") or {}).get(k, "")
        elif k == "mx_ok":
            out[k] = {1: "yes", 0: "no"}.get(l.get("mx_ok"), "")
        elif k == "people":
            out[k] = "; ".join(f"{p['name']} ({p['role']})" if p.get("role") else p["name"]
                               for p in l.get("people") or [])
        elif k == "people_emails":
            out[k] = "; ".join(f"{p['email']}{' (guess)' if p.get('email_status') == 'guessed' else ''}"
                               for p in l.get("people") or [] if p.get("email"))
        else:
            v = l.get(k)
            out[k] = "" if v is None else v
    return out


def _export_rows(args):
    f = _filters(args)
    rows = store.iter_leads(f)
    if args.get("only_email") == "1":
        rows = [r for r in rows if r.get("email")]
    return rows


@app.get("/api/export")
def export():
    fmt = request.args.get("fmt", "csv")
    rows = _export_rows(request.args)
    stamp = time.strftime("%Y%m%d-%H%M")
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(write_only=False)
        ws = wb.active
        ws.title = "Leads"
        ws.append([h for _, h in EXPORT_COLS])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="4F46E5")
        for r in rows:
            f = _flat(r)
            ws.append([f[k] for k, _ in EXPORT_COLS])
        widths = {"name": 34, "email": 32, "emails": 40, "website": 34, "address": 40}
        for i, (k, _) in enumerate(EXPORT_COLS):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = widths.get(k, 16)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="leads-{stamp}.xlsx"'})
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in EXPORT_COLS])
    for r in rows:
        f = _flat(r)
        w.writerow([f[k] for k, _ in EXPORT_COLS])
    return Response("\ufeff" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="leads-{stamp}.csv"'})


# ------------------------------------------------------------------ MailBlaster
def _mailblaster_db():
    custom = store.setting("mailblaster_db")
    cands = [custom] if custom else []
    cands += [os.path.join(os.environ.get("LOCALAPPDATA", ""), "MailBlaster", "mailblaster.db"),
              r"D:\mailblaster\mailblaster.db"]
    found = [p for p in cands if p and os.path.isfile(p)]
    if custom and custom in found:
        return custom
    return max(found, key=os.path.getmtime) if found else ""


@app.get("/api/mailblaster")
def mailblaster_info():
    p = _mailblaster_db()
    if not p:
        return jsonify({"found": False})
    try:
        c = sqlite3.connect(p, timeout=15)
        groups = [{"id": r[0], "name": r[1], "n": r[2]} for r in c.execute(
            "SELECT l.id,l.name,COUNT(m.contact_id) FROM lists l LEFT JOIN list_members m ON m.list_id=l.id "
            "GROUP BY l.id ORDER BY l.name")]
        c.close()
    except sqlite3.Error as ex:
        return jsonify({"found": True, "path": p, "error": str(ex)})
    return jsonify({"found": True, "path": p, "groups": groups})


@app.post("/api/export/mailblaster")
def export_mailblaster():
    d = request.get_json(force=True) or {}
    group = (d.get("group") or "").strip() or "Lead Finder " + time.strftime("%Y-%m-%d")
    all_emails = bool(d.get("all_emails"))
    with_people = bool(d.get("with_people", True))
    ids = _resolve_ids(d)
    leads = [store.get_lead(i) for i in ids]
    leads = [l for l in leads if l and l.get("email")]
    path = _mailblaster_db()
    if not path:
        return jsonify({"error": "MailBlaster database not found. Set its path in Settings."}), 400
    c = sqlite3.connect(path, timeout=30)
    c.row_factory = sqlite3.Row
    added = updated = suppressed = 0
    try:
        cols = {r[1] for r in c.execute("PRAGMA table_info(contacts)")}
        lcols = {r[1] for r in c.execute("PRAGMA table_info(lists)")}
        row = c.execute("SELECT id FROM lists WHERE name=?", (group,)).fetchone()
        if row:
            gid = row[0]
        elif "color" in lcols:
            gid = c.execute("INSERT INTO lists(name,color,created_at) VALUES(?,?,?)",
                            (group, "#6366f1", time.time())).lastrowid
        else:
            gid = c.execute("INSERT INTO lists(name,created_at) VALUES(?,?)", (group, time.time())).lastrowid
        try:
            sup = {r[0]: r[1] for r in c.execute("SELECT email,reason FROM suppression")}
        except sqlite3.Error:
            sup = {}
        for l in leads:
            emails = (l["emails"] or [l["email"]]) if all_emails else [l["email"]]
            data = {k: v for k, v in {
                "company": l["name"], "website": l["website"], "phone": l["phone"], "city": l["city"],
                "country": l["country"], "category": l["category"] or l["niche"], "address": l["address"],
                "source": "Lead Finder"}.items() if v}
            targets = [(e, "", "", data) for e in emails[:5]]
            if with_people:  # decision makers become their own contacts, with a name to personalise with
                for p in l.get("people") or []:
                    ok = p.get("email_status") == "published" or p.get("email_check") == "valid"
                    if p.get("email") and ok:
                        first, _, last = p["name"].partition(" ")
                        targets.append((p["email"], first, last, dict(data, role=p.get("role", ""))))
            for em, first, last, pdata in targets:
                em = em.strip().lower()
                ex = c.execute("SELECT id,data FROM contacts WHERE email=?", (em,)).fetchone()
                if ex:
                    old = json.loads(ex["data"] or "{}")
                    old.update({k: v for k, v in pdata.items() if not old.get(k)})
                    c.execute("UPDATE contacts SET data=? WHERE id=?", (json.dumps(old), ex["id"]))
                    if first:
                        c.execute("UPDATE contacts SET first_name=?, last_name=? WHERE id=? AND first_name='",
                                  (first, last, ex["id"]))
                    cid = ex["id"]
                    updated += 1
                else:
                    reason = sup.get(em)
                    status = ("bounced" if reason == "bounce" else "unsubscribed") if reason else "active"
                    suppressed += 1 if reason else 0
                    if "status_reason" in cols:
                        cid = c.execute("INSERT INTO contacts(email,first_name,last_name,data,status,status_reason,"
                                        "created_at) VALUES(?,?,?,?,?,?,?)",
                                        (em, first, last, json.dumps(pdata), status, reason or "",
                                         time.time())).lastrowid
                    else:
                        cid = c.execute("INSERT INTO contacts(email,first_name,last_name,data,status,created_at) "
                                        "VALUES(?,?,?,?,?,?)",
                                        (em, first, last, json.dumps(pdata), status, time.time())).lastrowid
                    added += 1
                c.execute("INSERT OR IGNORE INTO list_members(list_id,contact_id) VALUES(?,?)", (gid, cid))
        c.commit()
    except sqlite3.Error as ex:
        c.rollback()
        return jsonify({"error": f"MailBlaster database error: {ex}"}), 500
    finally:
        c.close()
    store.mark_exported([l["id"] for l in leads])
    return jsonify({"group": group, "leads": len(leads), "added": added, "updated": updated,
                    "suppressed": suppressed, "skipped_no_email": len(ids) - len(leads)})


# ------------------------------------------------------------------ settings
@app.get("/api/settings")
def settings_get():
    s = store.get_settings()
    s["mailblaster_detected"] = _mailblaster_db()
    s["data_dir"] = app_paths.data_dir()
    return jsonify(s)


@app.post("/api/settings")
def settings_save():
    store.save_settings(request.get_json(force=True) or {})
    return jsonify(store.get_settings())


@app.post("/api/settings/test")
def settings_test():
    d = request.get_json(force=True) or {}
    which, key = d.get("which"), (d.get("key") or "").strip()
    import requests as rq
    try:
        if which == "google":
            r = rq.post("https://places.googleapis.com/v1/places:searchText", timeout=20,
                        headers={"X-Goog-Api-Key": key, "X-Goog-FieldMask": "places.id"},
                        json={"textQuery": "coffee in Paris", "pageSize": 1})
            ok = r.status_code == 200
            msg = "Key works." if ok else r.json().get("error", {}).get("message", r.text[:200])
        elif which == "yelp":
            r = rq.get("https://api.yelp.com/v3/businesses/search", timeout=20,
                       headers={"Authorization": f"Bearer {key}"}, params={"term": "coffee", "location": "Paris", "limit": 1})
            ok = r.status_code == 200
            msg = "Key works." if ok else r.json().get("error", {}).get("description", r.text[:200])
        elif which == "companies_house":
            r = rq.get(registries.GB_API + "/search/companies", timeout=20, params={"q": "tesco", "items_per_page": 1},
                       headers=registries._gb_auth(key))
            ok = r.status_code == 200
            msg = "Key works." if ok else f"Companies House said HTTP {r.status_code} — check the key (REST API key)."
        elif which == "brave":
            r = rq.get("https://api.search.brave.com/res/v1/web/search", timeout=20, params={"q": "bakery paris"},
                       headers={"X-Subscription-Token": key, "Accept": "application/json"})
            ok = r.status_code == 200
            msg = "Key works." if ok else f"Brave Search said HTTP {r.status_code}: {r.text[:150]}"
        elif which == "hf":
            h = {"Authorization": f"Bearer {key}"}
            who = rq.get("https://huggingface.co/api/whoami-v2", headers=h, timeout=20)
            if who.status_code != 200:
                ok, msg = False, "Hugging Face rejected this token."
            else:
                r = rq.get(f"https://huggingface.co/datasets/{overture.FSQ_REPO}/resolve/main/README.md",
                           headers=h, timeout=20)
                ok = r.status_code == 200
                msg = "Token works and Foursquare access is granted." if ok else (
                    "Token works, but open huggingface.co/datasets/foursquare/fsq-os-places and accept the "
                    "terms first (approval is instant).")
        else:
            return jsonify({"ok": False, "msg": "unknown"}), 400
    except Exception as ex:
        ok, msg = False, str(ex)
    return jsonify({"ok": ok, "msg": msg})


@app.get("/api/ping")
def ping():
    return jsonify({"app": "lead-finder", "version": VERSION})


# ------------------------------------------------------------------ main
VERSION = "1.1.0"
DEFAULT_PORT = 5070


def _pick_port(start):
    for p in range(start, start + 50):
        with socket.socket() as s:
            if s.connect_ex(("127.0.0.1", p)) != 0:
                return p
    return start


def _already_running(port) -> bool:
    """A second launch (Start-menu click) just reopens the running copy in the browser."""
    import urllib.request
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/ping", timeout=1.5) as r:
            return json.loads(r.read().decode("utf-8")).get("app") == "lead-finder"
    except Exception:
        return False


def _redirect_logs():
    """pythonw has no console: send output to a log file in the data folder."""
    try:
        sys.stdout = sys.stderr = open(os.path.join(app_paths.data_dir(), "leadfinder.log"), "a",
                                       encoding="utf-8", buffering=1)
    except OSError:
        sys.stdout = sys.stderr = open(os.devnull, "w")


if __name__ == "__main__":
    if sys.stdout is None or app_paths.FROZEN:
        _redirect_logs()
    open_browser = (app_paths.is_installed() or os.environ.get("LF_OPEN") == "1") and os.environ.get("LF_NO_BROWSER") != "1"
    if not os.environ.get("PORT") and _already_running(DEFAULT_PORT):
        if open_browser:
            webbrowser.open(f"http://127.0.0.1:{DEFAULT_PORT}")
        sys.exit(0)
    store.init()
    port = int(os.environ.get("PORT", "0")) or _pick_port(DEFAULT_PORT)
    url = f"http://127.0.0.1:{port}"
    for sid in store.INTERRUPTED:  # searches cut off by a close / crash / update carry on by themselves
        threading.Timer(3, engine.resume, args=(sid,)).start()
    print(f"\n  Lead Finder is running at {url}\n  (close this window to stop)\n")
    if open_browser:
        threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    app.run(host="127.0.0.1", port=port, threaded=True, debug=False)
